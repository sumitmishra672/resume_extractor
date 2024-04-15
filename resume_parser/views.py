from django.shortcuts import render,redirect
from .models import Resume
# from docx2txt import process
import os
import re
from docx import Document
from openpyxl import Workbook
from PyPDF2 import PdfReader
from django.http import FileResponse

# Create your views here.

def upload_resume(request):
    if request.method == 'POST':
            file=request.FILES.get('upload')
            ress=Resume(file=file)
            ress.save()
            return redirect('data_ext')
    else:
        return render(request, 'upload_resume.html')
def data_ext(request):
        res=Resume.objects.order_by('-id')[:1]
        for x in res:
            data=x.file
        print(data)
        input_file = os.path.join(f"{data}")  # Update with your input file path
        output_file = os.path.join("output",'output.xlsx')    # Output Excel file path

        if input_file.endswith('.docx'):
            doc = Document(input_file)
            text = '\n'.join([para.text for para in doc.paragraphs])
        elif input_file.endswith('.pdf'):
            with open(input_file, 'rb') as file:
                reader = PdfReader(file)
                text = ''
                for page in reader.pages:
                    text += page.extract_text()
        else:
            print('Unsupported file format.')
            exit()

        names = re.findall(r'\b[A-Za-z]+\s[A-Za-z]+\b', text)
        emails = re.findall(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', text)
        phones = re.findall(r'\b\d{3}[-.\s]?\d{3}[-.\s]?\d{4}\b', text)
        data = list(zip(names, emails, phones))

        wb = Workbook()
        ws = wb.active
        ws.append(['Name', 'Email', 'Phone'])
        for row in data:
            ws.append(row)

        wb.save(output_file)
        print(f'Data extracted and saved to {output_file}.')
        
        return render(request, 'upload_success.html')
def download(request):
     files=os.path.join('output/output.xlsx')
     fileOpen=open(files,'rb')
     return FileResponse(fileOpen)

    


